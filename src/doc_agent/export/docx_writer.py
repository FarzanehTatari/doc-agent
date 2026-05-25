"""Render a `Bundle` to a single .docx file.

Approach: tokenize each entry's markdown with markdown-it-py, then walk the
flat token list and emit python-docx elements. This avoids re-implementing a
markdown parser while keeping us in control of the docx side (styles, tables,
page breaks). Supports the subset of markdown the agent actually produces:

  headings (h1–h4), paragraphs, inline emphasis (bold, italic, code, link),
  bullet and numbered lists, tables, fenced code blocks, blockquotes, hr.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from doc_agent.export.bundle import Bundle


def write_docx(bundle: Bundle, out_path: Path | str) -> Path:
    """Write `bundle` to a single .docx file. Returns the path written."""
    from docx import Document
    from docx.shared import Pt
    from markdown_it import MarkdownIt

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    md = (
        MarkdownIt("commonmark", {"breaks": True, "html": False, "linkify": True})
        .enable("table")
        .enable("strikethrough")
    )

    doc = Document()
    # Tighten default paragraph style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # --- Cover page ---
    doc.add_heading(f"{bundle.model_name} — Documentation", level=0)
    p = doc.add_paragraph()
    p.add_run(
        f"{len(bundle.entries)} document(s) — generated "
        f"{datetime.now(timezone.utc).isoformat(timespec='seconds')}"
    ).italic = True

    # --- Table of contents ---
    doc.add_paragraph()
    doc.add_heading("Contents", level=1)
    for entry in bundle.entries:
        para = doc.add_paragraph(style="List Bullet")
        kind_run = para.add_run(f"[{entry.kind}] ")
        kind_run.bold = True
        para.add_run(entry.title)

    # --- Body ---
    for entry in bundle.entries:
        doc.add_page_break()
        tokens = md.parse(entry.content)
        _walk(tokens, doc)

    doc.save(out_path)
    return out_path


# =============================================================================
# Token walker — converts a flat list of markdown-it tokens to docx elements.
#
# markdown-it tokens come in matched pairs (heading_open / heading_close, etc.)
# with an `inline` token in between carrying the actual content. We track the
# current "open container" (paragraph, list, etc.) to know where to put content.
# =============================================================================
def _walk(tokens: list, doc: Any) -> None:
    """Walk markdown-it tokens, emitting docx elements into `doc`."""
    state = _WalkState(doc=doc)
    i = 0
    while i < len(tokens):
        t = tokens[i]
        if t.type == "heading_open":
            level = int(t.tag[1])  # 'h1' → 1
            inline = tokens[i + 1] if i + 1 < len(tokens) else None
            text = inline.content if inline and inline.type == "inline" else ""
            # python-docx accepts levels 0-9; map our markdown h1-h4 to docx 1-4
            doc.add_heading(text.strip(), level=min(max(level, 1), 4))
            i += 3  # heading_open, inline, heading_close
            continue
        if t.type == "paragraph_open":
            inline = tokens[i + 1] if i + 1 < len(tokens) else None
            if inline and inline.type == "inline":
                state.add_paragraph_from_inline(inline)
            i += 3
            continue
        if t.type == "hr":
            # Simple horizontal rule — a line of dashes is the most portable
            doc.add_paragraph("─" * 60)
            i += 1
            continue
        if t.type == "bullet_list_open":
            i = state.consume_list(tokens, i, ordered=False)
            continue
        if t.type == "ordered_list_open":
            i = state.consume_list(tokens, i, ordered=True)
            continue
        if t.type == "fence" or t.type == "code_block":
            state.add_code_block(t.content)
            i += 1
            continue
        if t.type == "blockquote_open":
            i = state.consume_blockquote(tokens, i)
            continue
        if t.type == "table_open":
            i = state.consume_table(tokens, i)
            continue
        # Unknown / closing token — just advance
        i += 1


class _WalkState:
    """Helper holding shared logic for the token walker."""

    def __init__(self, doc: Any):
        self.doc = doc

    # --- paragraph ---------------------------------------------------------
    def add_paragraph_from_inline(self, inline, style: str | None = None) -> None:
        para = self.doc.add_paragraph(style=style) if style else self.doc.add_paragraph()
        self._emit_inline_children(inline, para)

    def add_code_block(self, text: str) -> None:
        from docx.shared import Pt
        para = self.doc.add_paragraph()
        run = para.add_run(text.rstrip("\n"))
        run.font.name = "Consolas"
        run.font.size = Pt(10)

    # --- inline content (bold/italic/code/link/text) -----------------------
    def _emit_inline_children(self, inline, para) -> None:
        """Walk an `inline` token's children, emitting runs into `para`."""
        children = inline.children or []
        bold = italic = code = False
        link_target = None
        for c in children:
            t = c.type
            if t == "strong_open":   bold = True
            elif t == "strong_close": bold = False
            elif t == "em_open":      italic = True
            elif t == "em_close":     italic = False
            elif t == "code_inline":
                run = para.add_run(c.content)
                run.font.name = "Consolas"
                run.bold = bold
                run.italic = italic
            elif t == "link_open":
                link_target = next(
                    (v for k, v in (c.attrs or {}).items() if k == "href"), None
                )
            elif t == "link_close":
                link_target = None
            elif t == "softbreak" or t == "hardbreak":
                para.add_run("\n")
            elif t == "text":
                content = c.content
                if link_target:
                    content = f"{content} ({link_target})"
                run = para.add_run(content)
                run.bold = bold
                run.italic = italic

    # --- lists -------------------------------------------------------------
    def consume_list(self, tokens, start: int, *, ordered: bool) -> int:
        """Process a bullet_list or ordered_list block. Return the index after it."""
        style = "List Number" if ordered else "List Bullet"
        i = start + 1
        end_type = "ordered_list_close" if ordered else "bullet_list_close"
        while i < len(tokens) and tokens[i].type != end_type:
            t = tokens[i]
            if t.type == "list_item_open":
                # Walk the children until list_item_close, accumulating into one para
                para = self.doc.add_paragraph(style=style)
                j = i + 1
                while j < len(tokens) and tokens[j].type != "list_item_close":
                    inner = tokens[j]
                    if inner.type == "paragraph_open":
                        if j + 1 < len(tokens) and tokens[j + 1].type == "inline":
                            self._emit_inline_children(tokens[j + 1], para)
                        j += 3
                        continue
                    if inner.type == "bullet_list_open" or inner.type == "ordered_list_open":
                        # Nested list — just walk it (will produce extra paragraphs)
                        j = self.consume_list(
                            tokens, j, ordered=(inner.type == "ordered_list_open")
                        )
                        continue
                    j += 1
                i = j + 1  # skip list_item_close
                continue
            i += 1
        return i + 1  # skip end_type

    # --- blockquote --------------------------------------------------------
    def consume_blockquote(self, tokens, start: int) -> int:
        """Process a blockquote block. Return the index after it."""
        i = start + 1
        while i < len(tokens) and tokens[i].type != "blockquote_close":
            t = tokens[i]
            if t.type == "paragraph_open":
                inline = tokens[i + 1] if i + 1 < len(tokens) else None
                para = self.doc.add_paragraph(style="Intense Quote")
                if inline and inline.type == "inline":
                    self._emit_inline_children(inline, para)
                i += 3
                continue
            i += 1
        return i + 1

    # --- tables ------------------------------------------------------------
    def consume_table(self, tokens, start: int) -> int:
        """Build a docx Table from the token stream.

        Each cell's *inline token* (not just its raw text) is stored so we can
        route it through `_emit_inline_children` and pick up bold / italic /
        `code` / link formatting inside cells — otherwise backticks etc. would
        appear as literal text in Word.
        """
        rows_inline: list[list] = []   # rows of inline tokens (None for empty cells)
        i = start + 1
        while i < len(tokens) and tokens[i].type != "table_close":
            t = tokens[i]
            if t.type == "tr_open":
                cells_inline: list = []
                j = i + 1
                while j < len(tokens) and tokens[j].type != "tr_close":
                    tj = tokens[j]
                    if tj.type in ("th_open", "td_open"):
                        inline = (
                            tokens[j + 1]
                            if j + 1 < len(tokens) and tokens[j + 1].type == "inline"
                            else None
                        )
                        cells_inline.append(inline)
                        # th_open/td_open, inline, th_close/td_close → 3 tokens
                        j += 3
                        continue
                    j += 1
                rows_inline.append(cells_inline)
                i = j + 1
                continue
            i += 1

        if rows_inline:
            n_cols = max(len(r) for r in rows_inline)
            table = self.doc.add_table(rows=len(rows_inline), cols=n_cols)
            try:
                table.style = "Light Grid"
            except KeyError:
                pass
            for ri, row in enumerate(rows_inline):
                for ci, inline in enumerate(row):
                    cell = table.rows[ri].cells[ci]
                    # python-docx creates one empty paragraph per cell by default.
                    # Reuse it: clear text and emit inline runs into it.
                    para = cell.paragraphs[0]
                    para.text = ""
                    if inline is not None:
                        self._emit_inline_children(inline, para)
                    if ri == 0:
                        # Header row — bold every run that's in the cell now
                        for run in para.runs:
                            run.bold = True
        return i + 1

"""DOCX writer tests — requires python-docx + markdown-it-py installed."""

import pytest

pytest.importorskip("markdown_it")
pytest.importorskip("docx")

from doc_agent.export.bundle import Bundle, BundleEntry
from doc_agent.export.docx_writer import write_docx


def _bundle(entries):
    return Bundle(model_name="M", entries=entries)


def _entry(title, content, kind="autodoc"):
    from pathlib import Path
    return BundleEntry(source_path=Path(f"{title}.md"),
                       content=content, title=title, kind=kind)


def _read_docx_text(path) -> str:
    """Concatenate all paragraph text from a .docx file."""
    from docx import Document
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def test_write_docx_produces_file(tmp_path):
    bundle = _bundle([_entry("A", "# A\n\nHello world.")])
    out = tmp_path / "out.docx"
    written = write_docx(bundle, out)
    assert written == out
    assert out.exists()
    txt = _read_docx_text(out)
    assert "A" in txt
    assert "Hello world." in txt


def test_docx_contains_cover_and_toc(tmp_path):
    bundle = _bundle([
        _entry("First",  "# First\n\nbody1", kind="autodoc"),
        _entry("Second", "# Second\n\nbody2", kind="sysreq"),
    ])
    out = tmp_path / "out.docx"
    write_docx(bundle, out)
    txt = _read_docx_text(out)
    assert "M — Documentation" in txt
    assert "Contents" in txt
    assert "[autodoc] First" in txt
    assert "[sysreq] Second" in txt


def test_docx_renders_inline_emphasis(tmp_path):
    bundle = _bundle([_entry("E",
        "# E\n\nThis has **bold**, *italic*, `code`, and a [link](https://example.com)."
    )])
    out = tmp_path / "e.docx"
    write_docx(bundle, out)
    txt = _read_docx_text(out)
    assert "bold" in txt
    assert "italic" in txt
    assert "code" in txt
    # Links are emitted as "text (url)"
    assert "https://example.com" in txt


def test_docx_renders_lists(tmp_path):
    bundle = _bundle([_entry("L",
        "# L\n\n- alpha\n- beta\n- gamma\n"
    )])
    out = tmp_path / "l.docx"
    write_docx(bundle, out)
    txt = _read_docx_text(out)
    for item in ("alpha", "beta", "gamma"):
        assert item in txt


def test_docx_renders_tables(tmp_path):
    bundle = _bundle([_entry("T",
        "# T\n\n| Name | Value |\n|---|---|\n| K_FOO | 0.5 |\n| K_BAR | 1.0 |\n"
    )])
    out = tmp_path / "t.docx"
    write_docx(bundle, out)
    from docx import Document
    doc = Document(str(out))
    assert len(doc.tables) >= 1
    tbl = doc.tables[0]
    assert tbl.rows[0].cells[0].text == "Name"
    assert any(r.cells[0].text == "K_FOO" for r in tbl.rows)


def test_docx_renders_code_block(tmp_path):
    bundle = _bundle([_entry("C",
        "# C\n\n```\nOut1 = K_VSE_FILT_TC * In1\n```\n"
    )])
    out = tmp_path / "c.docx"
    write_docx(bundle, out)
    txt = _read_docx_text(out)
    assert "Out1 = K_VSE_FILT_TC * In1" in txt
